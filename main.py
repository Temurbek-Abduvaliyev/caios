import os, io, re, json, html, time, asyncio, secrets, logging
from contextlib import asynccontextmanager

import asyncpg, httpx
from fastapi import FastAPI, Request, HTTPException
from aiogram import Bot, Dispatcher, Router, F
from aiogram.client.default import DefaultBotProperties
from aiogram.enums import ParseMode
from aiogram.filters import CommandStart, Command, CommandObject
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage
from aiogram.types import (Update, Message, CallbackQuery, InlineKeyboardMarkup,
                           InlineKeyboardButton, BufferedInputFile)
from docx import Document
from pptx import Presentation
from pptx.util import Pt
import os

PORT = int(os.environ.get("PORT", 5000))

logging.basicConfig(level=logging.INFO)
log = logging.getLogger("bot")

# ---------- Sozlamalar (Render Environment ga yoziladi) ----------
BOT_TOKEN = os.environ["BOT_TOKEN"]
DATABASE_URL = os.environ["DATABASE_URL"]
GROQ_KEYS = [k.strip() for k in os.environ["GROQ_KEYS"].split(",") if k.strip()]
MODEL = os.getenv("GROQ_MODEL", "openai/gpt-oss-20b")
REASONING = os.getenv("REASONING", "low")
ADMIN_IDS = {int(x) for x in os.getenv("ADMIN_IDS", "").split(",") if x.strip()}
ADMIN_CONTACT = os.getenv("ADMIN_CONTACT", "admin")
BASE_URL = os.getenv("WEBHOOK_URL") or os.getenv("RENDER_EXTERNAL_URL", "")
SECRET = os.getenv("WEBHOOK_SECRET", "change-me-please")
WORKERS = int(os.getenv("WORKERS", "2"))
GROQ_URL = "https://api.groq.com/openai/v1/chat/completions"

# ---------- Hujjat turlari ----------
DOC_SYS = ("Sen akademik yozuvchi yordamchisan. Faqat o'zbek tilida (lotin yozuvi), grammatik to'g'ri va rasmiy "
           "uslubda yoz. Markdown ishlat: # sarlavha, ## bo'lim, - ro'yxat. Jadval ishlatma. "
           "Kirish yoki izoh gap qo'shma, faqat hujjat matnini ber.")
SLIDE_SYS = ("Sen taqdimot yaratuvchisan. O'zbek tilida (lotin yozuvi) yoz. Faqat JSON qaytar, boshqa matn yo'q: "
             '{"title":"...","subtitle":"...","slides":[{"title":"...","bullets":["...","..."]}]}. '
             "8-10 ta slayd, har birida 3-5 band, har band 14 so'zdan oshmasin.")
DOCS = {
    "esse": ("Esse", "Esse yoz, taxminan 700 so'z: kirish, 3 ta asosiy fikr (dalil bilan), xulosa."),
    "referat": ("Referat", "Referat yoz, taxminan 1800 so'z: reja, kirish, 3-4 bob, xulosa, foydalanilgan adabiyotlar. "
                "Mavjudligiga ishonching komil bo'lmagan kitob yoki havolalarni to'qima."),
    "maqola": ("Maqola", "Ilmiy-ommabop maqola yoz, taxminan 1000 so'z: sarlavha, annotatsiya, kirish, asosiy qism, xulosa."),
    "mustaqil": ("Mustaqil ish", "Mustaqil ish yoz, taxminan 1500 so'z: reja, kirish, bo'limlarga bo'lingan asosiy qism, xulosa."),
    "tezis": ("Tezis", "Konferensiya tezisi yoz, 300-400 so'z: maqsad, usul, asosiy natijalar, xulosa."),
    "glossariy": ("Glossariy", "25 ta atamadan iborat glossariy yoz, alifbo tartibida, har biri '- Atama: qisqa izoh' shaklida."),
    "test": ("Test", "20 ta test savoli yoz. Har savolda A, B, C, D variantlar (har biri alohida qatorda). "
             "Oxirida '## Javoblar' bo'limida to'g'ri javoblar."),
    "quiz": ("Quiz", "10 ta quiz savoli yoz: har birida 4 variant. Oxirida '## Javoblar' bo'limida to'g'ri javob va 1 jumlali izoh."),
}
LABELS = {**{k: v[0] for k, v in DOCS.items()}, "slayd": "Slayd"}

# ---------- Groq kalit almashtirgich ----------
class KeyPool:
    def __init__(self, keys):
        self.keys, self.i, self.cool = keys, 0, {k: 0.0 for k in keys}
    def pick(self):
        now = time.time()
        for _ in range(len(self.keys)):
            k = self.keys[self.i % len(self.keys)]
            self.i += 1
            if self.cool[k] <= now:
                return k
        return None
    def rest(self, k, sec):
        self.cool[k] = time.time() + sec

pool = KeyPool(GROQ_KEYS)

async def ask(system, user, max_tokens=6000):
    async with httpx.AsyncClient(timeout=120) as http:
        for _ in range(len(GROQ_KEYS) * 2 + 3):
            k = pool.pick()
            if not k:
                await asyncio.sleep(5)
                continue
            try:
                r = await http.post(GROQ_URL, headers={"Authorization": f"Bearer {k}"}, json={
                    "model": MODEL, "reasoning_effort": REASONING, "max_tokens": max_tokens,
                    "messages": [{"role": "system", "content": system}, {"role": "user", "content": user}]})
            except httpx.HTTPError:
                await asyncio.sleep(2)
                continue
            if r.status_code == 200:
                text = (r.json()["choices"][0]["message"].get("content") or "").strip()
                if text:
                    return text
                continue
            if r.status_code == 429:
                pool.rest(k, float(r.headers.get("retry-after", 30)) + 1)
            elif r.status_code in (401, 403):
                pool.rest(k, 3600)
                log.warning("Kalit rad etildi (401/403)")
            elif r.status_code >= 500:
                await asyncio.sleep(3)
            else:
                raise RuntimeError(f"AI xatosi {r.status_code}: {r.text[:200]}")
    raise RuntimeError("AI hozir band")

# ---------- Fayl yaratish ----------
def make_docx(text, title):
    d = Document()
    d.add_heading(title, 0)
    for line in text.splitlines():
        s = line.strip().replace("**", "").replace("`", "")
        if not s:
            continue
        if s.startswith("### "): d.add_heading(s[4:], 3)
        elif s.startswith("## "): d.add_heading(s[3:], 2)
        elif s.startswith("# "): d.add_heading(s[2:], 1)
        elif s[:2] in ("- ", "* "): d.add_paragraph(s[2:], style="List Bullet")
        else: d.add_paragraph(s)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()

def make_pptx(raw):
    j = json.loads(raw[raw.index("{"): raw.rindex("}") + 1])
    p = Presentation()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = j.get("title", "Taqdimot")
    s.placeholders[1].text = j.get("subtitle", "")
    for sl in j.get("slides", []):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = sl.get("title", "")
        tf = s.placeholders[1].text_frame
        for n, b in enumerate(sl.get("bullets", [])):
            para = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            para.text = str(b)
            para.font.size = Pt(20)
    buf = io.BytesIO()
    p.save(buf)
    return buf.getvalue()

async def build(kind, topic):
    safe = re.sub(r"[^\w\-]+", "_", topic)[:40] or "hujjat"
    if kind == "slayd":
        return make_pptx(await ask(SLIDE_SYS, f"Mavzu: {topic}")), f"slayd_{safe}.pptx"
    text = await ask(DOC_SYS, f"{DOCS[kind][1]}\nMavzu: {topic}")
    return make_docx(text, topic), f"{kind}_{safe}.docx"

# ---------- Ma'lumotlar bazasi ----------
SCHEMA = """
CREATE TABLE IF NOT EXISTS users(id BIGINT PRIMARY KEY, name TEXT, credits INT NOT NULL DEFAULT 0,
  done INT NOT NULL DEFAULT 0, joined TIMESTAMPTZ NOT NULL DEFAULT now());
CREATE TABLE IF NOT EXISTS keys(code TEXT PRIMARY KEY, created TIMESTAMPTZ NOT NULL DEFAULT now(),
  used_by BIGINT, used_at TIMESTAMPTZ, revoked BOOLEAN NOT NULL DEFAULT FALSE);
"""
db: asyncpg.Pool = None
ALPH = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"
KEY_RE = r"^[A-Za-z0-9]{4}-[A-Za-z0-9]{4}-[A-Za-z0-9]{4}$"

def new_code():
    return "-".join("".join(secrets.choice(ALPH) for _ in range(4)) for _ in range(3))

async def make_keys(n):
    codes = [new_code() for _ in range(n)]
    await db.executemany("INSERT INTO keys(code) VALUES($1) ON CONFLICT DO NOTHING", [(c,) for c in codes])
    return codes

async def redeem(uid, code):
    async with db.acquire() as c, c.transaction():
        row = await c.fetchrow("UPDATE keys SET used_by=$1, used_at=now() "
                               "WHERE code=$2 AND used_by IS NULL AND NOT revoked RETURNING code", uid, code)
        if not row:
            return False
        await c.execute("UPDATE users SET credits=credits+1 WHERE id=$1", uid)
        return True

async def take_credit(uid):
    return await db.fetchval("UPDATE users SET credits=credits-1 WHERE id=$1 AND credits>0 RETURNING credits", uid)

async def credits(uid):
    return await db.fetchval("SELECT credits FROM users WHERE id=$1", uid) or 0

async def ensure_user(u):
    await db.execute("INSERT INTO users(id,name) VALUES($1,$2) ON CONFLICT (id) DO UPDATE SET name=$2", u.id, u.full_name)

# ---------- Bot ----------
bot = Bot(BOT_TOKEN, default=DefaultBotProperties(parse_mode=ParseMode.HTML))
dp = Dispatcher(storage=MemoryStorage())
r = Router()
dp.include_router(r)

class Flow(StatesGroup):
    topic = State()
    bmi = State()

queue: asyncio.Queue = asyncio.Queue()
busy: set = set()
fails: dict = {}

def menu_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📄 Hujjat yaratish", callback_data="m:docs")],
        [InlineKeyboardButton(text="⚖️ BMI (bepul)", callback_data="m:bmi")],
        [InlineKeyboardButton(text="🔑 Kalit kiritish", callback_data="m:key"),
         InlineKeyboardButton(text="👤 Balans", callback_data="m:bal")]])

def types_kb():
    items = list(LABELS.items())
    rows = [[InlineKeyboardButton(text=v, callback_data=f"t:{k}") for k, v in items[i:i + 2]]
            for i in range(0, len(items), 2)]
    return InlineKeyboardMarkup(inline_keyboard=rows)

NO_KEY = f"🔑 Sizda kalit yo'q. Kalit olish uchun {html.escape(ADMIN_CONTACT)} ga murojaat qiling, so'ng kalitni shu yerga yuboring."

@r.message(CommandStart())
async def start(m: Message, state: FSMContext):
    await state.clear()
    await ensure_user(m.from_user)
    await m.answer(f"Assalomu alaykum! Men hujjat yaratuvchi botman.\nBalans: <b>{await credits(m.from_user.id)}</b> ta so'rov.\n"
                   "Har bir kalit 1 ta hujjatga yetadi.", reply_markup=menu_kb())

@r.message(F.text.regexp(KEY_RE))
async def key_entered(m: Message):
    await ensure_user(m.from_user)
    uid = m.from_user.id
    cnt, t0 = fails.get(uid, (0, 0))
    if time.time() - t0 > 600:
        cnt = 0
    if cnt >= 5:
        return await m.answer("Juda ko'p noto'g'ri urinish. 10 daqiqadan keyin qayta urinib ko'ring.")
    if await redeem(uid, m.text.strip().upper()):
        fails.pop(uid, None)
        await m.answer(f"✅ Kalit qabul qilindi. Balans: <b>{await credits(uid)}</b> ta so'rov.", reply_markup=menu_kb())
    else:
        fails[uid] = (cnt + 1, t0 if cnt else time.time())
        await m.answer("❌ Kalit noto'g'ri yoki avval ishlatilgan.")

@r.callback_query(F.data.startswith("m:"))
async def menu(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    await ensure_user(cb.from_user)
    act = cb.data[2:]
    if act == "docs":
        await cb.message.answer("Hujjat turini tanlang:", reply_markup=types_kb())
    elif act == "bmi":
        await state.set_state(Flow.bmi)
        await cb.message.answer("Vazn (kg) va bo'yni (sm) yozing. Masalan: <code>70 175</code>")
    elif act == "key":
        await cb.message.answer("Kalitni yuboring (ko'rinishi: <code>XXXX-XXXX-XXXX</code>).")
    elif act == "bal":
        await cb.message.answer(f"Balans: <b>{await credits(cb.from_user.id)}</b> ta so'rov.")

@r.callback_query(F.data.startswith("t:"))
async def pick_type(cb: CallbackQuery, state: FSMContext):
    await cb.answer()
    kind = cb.data[2:]
    if kind not in LABELS:
        return
    if await credits(cb.from_user.id) < 1:
        return await cb.message.answer(NO_KEY)
    await state.set_state(Flow.topic)
    await state.update_data(kind=kind)
    await cb.message.answer(f"<b>{LABELS[kind]}</b> uchun mavzuni yozing:")

@r.message(Flow.topic, F.text)
async def got_topic(m: Message, state: FSMContext):
    uid = m.from_user.id
    kind = (await state.get_data()).get("kind")
    await state.clear()
    if uid in busy:
        return await m.answer("Avvalgi ishingiz tayyorlanmoqda, kuting.")
    if queue.qsize() >= 30:
        return await m.answer("Hozir navbat to'la, birozdan keyin urinib ko'ring.")
    if await take_credit(uid) is None:
        return await m.answer(NO_KEY)
    busy.add(uid)
    await queue.put((uid, m.chat.id, kind, m.text.strip()[:300]))
    await m.answer(f"⏳ Qabul qilindi. Navbatda: {queue.qsize()}. Tayyor bo'lgach fayl yuboriladi.")

@r.message(Flow.bmi, F.text)
async def got_bmi(m: Message, state: FSMContext):
    nums = re.findall(r"\d+(?:[.,]\d+)?", m.text)
    if len(nums) != 2:
        return await m.answer("Ikkita son yozing. Masalan: <code>70 175</code>")
    kg, cm = (float(x.replace(",", ".")) for x in nums)
    if not (20 <= kg <= 300 and 100 <= cm <= 250):
        return await m.answer("Qiymatlar noto'g'ri. Vazn 20-300 kg, bo'y 100-250 sm bo'lsin.")
    bmi = kg / (cm / 100) ** 2
    cat = "Kam vazn" if bmi < 18.5 else "Normal" if bmi < 25 else "Ortiqcha vazn" if bmi < 30 else "Semizlik"
    await state.clear()
    await m.answer(f"BMI: <b>{bmi:.1f}</b> — {cat}\n<i>Bu tibbiy tashxis emas.</i>", reply_markup=menu_kb())

# ---------- Admin panel ----------
def admin_kb():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text=f"➕ {n} ta kalit", callback_data=f"a:gen:{n}") for n in (1, 5, 10)],
        [InlineKeyboardButton(text="📋 Ishlatilmagan kalitlar", callback_data="a:list")],
        [InlineKeyboardButton(text="📊 Statistika", callback_data="a:stats")]])

@r.message(Command("admin"), F.from_user.id.in_(ADMIN_IDS))
async def admin(m: Message):
    await m.answer("🛠 <b>Admin panel</b>\nKalitni bekor qilish: <code>/revoke XXXX-XXXX-XXXX</code>", reply_markup=admin_kb())

@r.message(Command("revoke"), F.from_user.id.in_(ADMIN_IDS))
async def revoke(m: Message, command: CommandObject):
    code = (command.args or "").strip().upper()
    ok = await db.fetchval("UPDATE keys SET revoked=true WHERE code=$1 AND used_by IS NULL RETURNING code", code)
    await m.answer("🗑 Bekor qilindi." if ok else "Topilmadi yoki allaqachon ishlatilgan.")

@r.callback_query(F.data.startswith("a:"), F.from_user.id.in_(ADMIN_IDS))
async def admin_cb(cb: CallbackQuery):
    await cb.answer()
    p = cb.data.split(":")
    if p[1] == "gen":
        codes = await make_keys(int(p[2]))
        await cb.message.answer("🔑 Yangi kalitlar:\n" + "\n".join(f"<code>{c}</code>" for c in codes))
    elif p[1] == "list":
        rows = await db.fetch("SELECT code FROM keys WHERE used_by IS NULL AND NOT revoked ORDER BY created DESC LIMIT 30")
        await cb.message.answer("📋 Ishlatilmagan kalitlar:\n" + ("\n".join(f"<code>{x['code']}</code>" for x in rows) or "Yo'q"))
    elif p[1] == "stats":
        u = await db.fetchrow("SELECT count(*) n, coalesce(sum(done),0) d, coalesce(sum(credits),0) c FROM users")
        k = await db.fetchrow("SELECT count(*) t, count(used_by) used, "
                              "count(*) FILTER (WHERE used_by IS NULL AND NOT revoked) free FROM keys")
        await cb.message.answer(f"📊 Foydalanuvchilar: {u['n']}\nYaratilgan hujjatlar: {u['d']}\n"
                                f"Ishlatilmagan balans: {u['c']}\nKalitlar: jami {k['t']}, ishlatilgan {k['used']}, bo'sh {k['free']}")

# ---------- Navbat ishchilari ----------
async def worker():
    while True:
        uid, chat, kind, topic = await queue.get()
        ok = False
        try:
            data, name = await build(kind, topic)
            await bot.send_document(chat, BufferedInputFile(data, filename=name), caption="✅ Tayyor!")
            ok = True
        except Exception:
            log.exception("Ish bajarilmadi")
            await db.execute("UPDATE users SET credits=credits+1 WHERE id=$1", uid)
            try:
                await bot.send_message(chat, "❌ Xatolik yuz berdi. Kalitingiz qaytarildi, birozdan keyin qayta urinib ko'ring.")
            except Exception:
                pass
        finally:
            busy.discard(uid)
            queue.task_done()
        if ok:
            await db.execute("UPDATE users SET done=done+1 WHERE id=$1", uid)

# ---------- Web server (Render) ----------
tasks = []

@asynccontextmanager
async def lifespan(app):
    global db
    db = await asyncpg.create_pool(DATABASE_URL, min_size=1, max_size=3)
    await db.execute(SCHEMA)
    tasks.extend(asyncio.create_task(worker()) for _ in range(WORKERS))
    if BASE_URL:
        await bot.set_webhook(f"{BASE_URL}/webhook", secret_token=SECRET, drop_pending_updates=True)
    yield
    await db.close()
    await bot.session.close()

app = FastAPI(lifespan=lifespan)

@app.get("/health")
async def health():
    return {"ok": True}

@app.post("/webhook")
async def webhook(req: Request):
    if req.headers.get("x-telegram-bot-api-secret-token") != SECRET:
        raise HTTPException(403)
    await dp.feed_webhook_update(bot, Update.model_validate(await req.json(), context={"bot": bot}))
    return {"ok": True}
